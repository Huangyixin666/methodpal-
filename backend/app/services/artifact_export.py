"""Real binary exports from saved, owner-authorized artifacts.

No remote images, HTML execution or caller-supplied filesystem paths are used.
"""
import io
import os
import re
import threading
from pathlib import Path
from xml.sax.saxutils import escape

from fastapi import HTTPException

PLOT_LOCK = threading.RLock()

def clean(value):
    return re.sub(r'[\x00-\x08\x0b\x0c\x0e-\x1f]', '', str(value or ''))

def render_chart(chart, fmt='png'):
    if fmt not in ('png', 'svg'):
        raise HTTPException(422, '图表仅支持 PNG 和 SVG')
    import matplotlib
    matplotlib.use('Agg')
    from matplotlib import pyplot as plt, font_manager
    import numpy as np
    with PLOT_LOCK:
        font_path = next((p for p in [os.getenv('METHODPAL_CHART_FONT', ''),
            '/usr/share/fonts/opentype/noto/NotoSansCJK-Regular.ttc',
            'C:/Windows/Fonts/msyh.ttc'] if p and Path(p).is_file()), None)
        font = font_manager.FontProperties(fname=font_path) if font_path else None
        with matplotlib.rc_context({'axes.unicode_minus': False, 'font.size': 10, 'svg.fonttype': 'path'}):
            fig, ax = plt.subplots(figsize=(9, 5.6), constrained_layout=True)
            try:
                matrix = chart.get('matrix') or []
                rows = chart.get('data') or []
                ctype = chart.get('chart_type', '')
                if matrix:
                    data = np.array([[float(v) if v is not None else np.nan for v in row] for row in matrix])
                    image = ax.imshow(data, cmap='RdBu_r' if 'correlation' in ctype else 'Blues',
                        vmin=-1 if 'correlation' in ctype else None, vmax=1 if 'correlation' in ctype else None)
                    ax.set_xticks(range(len(chart.get('column_labels') or [])), chart.get('column_labels') or [], rotation=45, ha='right', fontproperties=font)
                    ax.set_yticks(range(len(chart.get('row_labels') or [])), chart.get('row_labels') or [], fontproperties=font)
                    fig.colorbar(image, ax=ax, shrink=.8)
                    if data.size <= 100:
                        for i in range(data.shape[0]):
                            for j in range(data.shape[1]):
                                if np.isfinite(data[i, j]):
                                    ax.text(j, i, f'{data[i,j]:.2g}', ha='center', va='center', fontsize=8)
                elif 'scatter' in ctype:
                    xk = chart.get('x_key') or 'y_true'
                    yk = chart.get('y_key') or 'y_pred'
                    points = [(r.get(xk), r.get(yk)) for r in rows if isinstance(r.get(xk), (int, float)) and isinstance(r.get(yk), (int, float))]
                    if points:
                        xs, ys = zip(*points)
                        ax.scatter(xs, ys, alpha=.7, color='#6554d9')
                        low, high = min(min(xs), min(ys)), max(max(xs), max(ys))
                        ax.plot([low, high], [low, high], '--', color='#888888')
                    ax.set_xlabel(chart.get('x_label') or xk, fontproperties=font)
                    ax.set_ylabel(chart.get('y_label') or yk, fontproperties=font)
                else:
                    xk = chart.get('x_key') or 'variable'
                    series = chart.get('series') or []
                    keys = [s.get('key') or s.get('dataKey') for s in series if s.get('key') or s.get('dataKey')]
                    if not keys:
                        keys = [chart.get('y_key') or next((k for r in rows for k, v in r.items() if isinstance(v, (int, float))), 'value')]
                    rows = rows[:30]
                    positions = np.arange(len(rows))
                    width = .8 / max(len(keys), 1)
                    labels = [clean(r.get(xk) or r.get('name') or r.get('model') or i + 1) for i, r in enumerate(rows)]
                    for index, key in enumerate(keys[:8]):
                        values = [float(r.get(key) or 0) if isinstance(r.get(key), (int, float)) else 0 for r in rows]
                        ax.barh(positions + index*width, values, height=width, label=clean(key))
                    ax.set_yticks(positions + width*(len(keys)-1)/2, labels, fontproperties=font)
                    ax.invert_yaxis()
                    if len(keys) > 1:
                        ax.legend(prop=font, fontsize=8)
                    ax.grid(axis='x', alpha=.2)
                ax.set_title(clean(chart.get('title')), fontproperties=font, fontsize=14, pad=15)
                buf = io.BytesIO()
                fig.savefig(buf, format=fmt, dpi=170, facecolor='white')
                return buf.getvalue()
            finally:
                plt.close(fig)

def markdown_blocks(content):
    from markdown_it import MarkdownIt
    tokens = MarkdownIt('commonmark').enable('table').parse(clean(content))
    i = 0
    while i < len(tokens):
        tok = tokens[i]
        if tok.type == 'table_open':
            rows, row = [], []
            i += 1
            while i < len(tokens) and tokens[i].type != 'table_close':
                if tokens[i].type == 'tr_open':
                    row = []
                elif tokens[i].type == 'inline':
                    row.append(tokens[i].content)
                elif tokens[i].type == 'tr_close':
                    rows.append(row)
                i += 1
            yield ('table', rows)
        elif tok.type == 'heading_open' and i + 1 < len(tokens):
            yield ('heading' + tok.tag[1:], tokens[i + 1].content)
            i += 1
        elif tok.type == 'inline':
            # Headings and tables were already handled above.
            yield ('paragraph', tok.content)
        elif tok.type in ('fence', 'code_block'):
            yield ('code', tok.content)
        i += 1

def plain(text):
    from markdown_it import MarkdownIt
    tokens = MarkdownIt('commonmark').parseInline(clean(text))
    return ''.join(child.content if child.type not in ('softbreak', 'hardbreak') else '\n'
        for token in tokens for child in token.children or [] if child.type in ('text', 'code_inline', 'softbreak', 'hardbreak'))

def report_export(report, fmt):
    if fmt not in ('markdown', 'docx', 'pdf'):
        raise HTTPException(422, '仅支持 Markdown、Word 和 PDF')
    content = report.get('draft_content', report.get('content') or '')
    provenance = report.get('source_snapshot') or {}
    version_text = '分析记录 ' + str(report.get('analysis_id') or '无') + '  数据版本 ' + str(provenance.get('dataset_revision', '旧版本'))
    if fmt == 'markdown':
        import base64
        exported = content + '\n\n## 分析版本\n\n' + version_text + '\n'
        for chart in report.get('charts') or []:
            image = base64.b64encode(render_chart(chart)).decode()
            exported += '\n\n![' + clean(chart.get('title')).replace(']', '') + '](data:image/png;base64,' + image + ')\n'
        return exported.encode('utf-8'), 'text/markdown; charset=utf-8'
    blocks = list(markdown_blocks(content))
    if fmt == 'docx':
        from docx import Document
        from docx.shared import Inches, Pt, RGBColor
        from docx.oxml import OxmlElement
        from docx.oxml.ns import qn
        from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
        doc = Document()
        # The bundled Word template may carry a colored rule on its title.
        for border in list(doc.styles.element.iter(qn('w:pBdr'))):
            border.getparent().remove(border)
        section = doc.sections[0]
        section.page_width, section.page_height = Inches(8.5), Inches(11)
        section.top_margin = section.bottom_margin = Inches(.75)
        section.left_margin = section.right_margin = Inches(.8)
        for name in ('Normal', 'Title', 'Heading 1', 'Heading 2', 'Heading 3'):
            style = doc.styles[name]
            style.font.name = 'Noto Sans CJK SC'
            style.font.color.rgb = RGBColor(0, 0, 0)
            style.element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'), 'Noto Sans CJK SC')
        doc.styles['Normal'].font.size = Pt(11)
        doc.styles['Normal'].paragraph_format.space_after = Pt(7)
        for index, (kind, text) in enumerate(blocks):
            if kind == 'table':
                if not text:
                    continue
                cols = max(len(r) for r in text)
                # Split wide matrices into readable panels rather than tiny type.
                for start in [0] + list(range(6, cols, 5)):
                    positions = list(range(min(cols, 6))) if start == 0 else [0] + list(range(start, min(cols, start + 5)))
                    table = doc.add_table(rows=0, cols=len(positions))
                    table.style = 'Table Grid'
                    table.alignment = WD_TABLE_ALIGNMENT.CENTER
                    for ri, values in enumerate(text):
                        cells = table.add_row().cells
                        for ci, pos in enumerate(positions):
                            cells[ci].text = plain(values[pos] if pos < len(values) else '')
                            cells[ci].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
                            tcpr = cells[ci]._tc.get_or_add_tcPr()
                            shade = OxmlElement('w:shd')
                            shade.set(qn('w:fill'), 'E8EDF3' if ri == 0 else 'FFFFFF' if ri % 2 else 'F6F8FA')
                            tcpr.append(shade)
                            borders = OxmlElement('w:tcBorders')
                            for side in ('top', 'left', 'bottom', 'right'):
                                edge = OxmlElement('w:' + side)
                                edge.set(qn('w:val'), 'single')
                                edge.set(qn('w:sz'), '4')
                                edge.set(qn('w:color'), 'D9D9D9')
                                borders.append(edge)
                            tcpr.append(borders)
                            margins = OxmlElement('w:tcMar')
                            for side in ('top', 'left', 'bottom', 'right'):
                                margin = OxmlElement('w:' + side)
                                margin.set(qn('w:w'), '90')
                                margin.set(qn('w:type'), 'dxa')
                                margins.append(margin)
                            tcpr.append(margins)
                        if ri == 0:
                            repeat = OxmlElement('w:tblHeader')
                            table.rows[0]._tr.get_or_add_trPr().append(repeat)
                    doc.add_paragraph()
            elif kind.startswith('heading'):
                level = min(int(kind[-1]), 3)
                doc.add_paragraph(plain(text), 'Title' if index == 0 and level == 1 else 'Heading ' + str(level))
            else:
                doc.add_paragraph(plain(text))
        doc.add_heading('分析版本', 1)
        doc.add_paragraph(version_text)
        for chart in report.get('charts') or []:
            p = doc.add_paragraph(clean(chart.get('title')), 'Heading 2')
            p.paragraph_format.keep_with_next = True
            doc.add_picture(io.BytesIO(render_chart(chart)), width=Inches(6.6))
            for note in chart.get('notes') or []:
                doc.add_paragraph(clean(note))
        out = io.BytesIO()
        doc.save(out)
        return out.getvalue(), 'application/vnd.openxmlformats-officedocument.wordprocessingml.document'
    from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle, Image, KeepTogether
    from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
    from reportlab.lib import colors
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.cidfonts import UnicodeCIDFont
    from reportlab.lib.pagesizes import letter
    pdfmetrics.registerFont(UnicodeCIDFont('STSong-Light'))
    styles = getSampleStyleSheet()
    for name in ('Normal', 'Title', 'Heading1', 'Heading2', 'Heading3'):
        styles[name].fontName = 'STSong-Light'
        styles[name].textColor = colors.black
        styles[name].wordWrap = 'CJK'
        styles[name].leading = styles[name].fontSize * 1.5
    styles['Normal'].fontSize, styles['Normal'].leading = 11, 16
    cell_style = ParagraphStyle('Cell', parent=styles['Normal'], fontSize=8, leading=12)
    flow = []
    def pdf_text(text):
        # Standard Chinese CID fonts do not reliably contain superscript digits.
        return escape(plain(text).replace('²', '^2').replace('³', '^3'))
    def para(text, style='Normal'):
        return Paragraph(pdf_text(text).replace('\n', '<br/>'), styles[style])
    for index, (kind, text) in enumerate(blocks):
        if kind == 'table':
            if not text:
                continue
            cols = max(len(row) for row in text)
            for start in range(0, cols, 5):
                positions = list(range(start, min(cols, start + 5))) if start == 0 else [0] + list(range(start, min(cols, start + 5)))
                data = [[Paragraph(pdf_text(row[pos] if pos < len(row) else ''), cell_style) for pos in positions] for row in text]
                table = Table(data, colWidths=[475 / len(positions)] * len(positions), repeatRows=1, hAlign='LEFT')
                table.setStyle(TableStyle([('BACKGROUND', (0,0), (-1,0), colors.HexColor('#E8EDF3')),
                    ('GRID', (0,0), (-1,-1), .4, colors.HexColor('#D9D9D9')), ('VALIGN',(0,0),(-1,-1),'MIDDLE'),
                    ('LEFTPADDING',(0,0),(-1,-1),6), ('RIGHTPADDING',(0,0),(-1,-1),6),
                    ('TOPPADDING',(0,0),(-1,-1),6), ('BOTTOMPADDING',(0,0),(-1,-1),6)]))
                flow.extend([table, Spacer(1, 10)])
        elif kind.startswith('heading'):
            level = min(int(kind[-1]), 3)
            flow.append(para(text, 'Title' if index == 0 and level == 1 else 'Heading' + str(level)))
        else:
            flow.extend([para(text), Spacer(1, 6)])
    flow.extend([para('分析版本', 'Heading1'), para(version_text)])
    for chart in report.get('charts') or []:
        flow.append(KeepTogether([para(chart.get('title'), 'Heading2'),
            Image(io.BytesIO(render_chart(chart)), width=475, height=296)]))
        flow.extend(para(note) for note in chart.get('notes') or [])
    out = io.BytesIO()
    SimpleDocTemplate(out, pagesize=letter, leftMargin=54, rightMargin=54,
        topMargin=48, bottomMargin=48, title=clean(report.get('title'))).build(flow)
    return out.getvalue(), 'application/pdf'
